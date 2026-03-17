import argparse
import json
import re
import subprocess
import tempfile
from pathlib import Path

import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from season_priors import _normalize_school_name, _slugify_school
from tournament_matchups import BRACKET_TEAM_ALIASES, default_field_source_url, fetch_bracket_pdf, require_pdftotext


DEFAULT_TOURNAMENT_DIR = Path("results/tournament")
DEFAULT_OUTPUT_DIR = Path("results/bracket")
REGION_NAMES = ["EAST", "WEST", "SOUTH", "MIDWEST"]
SEMIFINAL_PAIRINGS = [("EAST", "WEST"), ("SOUTH", "MIDWEST")]
SEED_SLOT_ORDER = [1, 16, 8, 9, 5, 12, 4, 13, 6, 11, 3, 14, 7, 10, 2, 15]
REGION_COLORS = {
    "EAST": "DDEBF7",
    "WEST": "E2F0D9",
    "SOUTH": "FCE4D6",
    "MIDWEST": "FFF2CC",
}
WINNER_FILL = PatternFill("solid", fgColor="D9EAD3")
LOSER_FILL = PatternFill("solid", fgColor="F4F4F4")
TITLE_FILL = PatternFill("solid", fgColor="1F4E78")
FINAL_FILL = PatternFill("solid", fgColor="D9D2E9")
THIN_SIDE = Side(style="thin", color="B7C0CC")
CELL_BORDER = Border(left=THIN_SIDE, right=THIN_SIDE, top=THIN_SIDE, bottom=THIN_SIDE)

PLAY_IN_NAME_ALIASES = {
    "how": "howard",
    "mia-oh": "miami-oh",
    "nc-st": "north-carolina-state",
    "pvamu": "prairie-view",
}

SEED_LINE_RE = re.compile(r"^(?P<seed>1[0-6]|[1-9])(?=\s|[A-Z])\s*(?P<body>.+?)\s*$")
RECORD_IN_LINE_RE = re.compile(r"^(?P<team>.+?)\s*\((?P<wins>\d+)-(?P<losses>\d+)\)(?:\s+.*)?$")
PLAY_IN_RE = re.compile(r"^(?P<left>.+?)\s+vs\s+(?P<right>.+?)$")
TEAM_WITH_RECORD_RE = re.compile(r"(?<!\d)(?P<seed>1[0-6]|[1-9])(?=\s|[A-Z])\s*(?P<body>.+?\(\d+\-\d+\))")


def parse_args():
    parser = argparse.ArgumentParser()
    parser.add_argument("--season-year", type=int, default=2026)
    parser.add_argument("--cutoff-date", default="2026-03-15")
    parser.add_argument("--field-source-url")
    parser.add_argument("--tournament-dir")
    parser.add_argument("--probability-matrix-path")
    parser.add_argument("--team-field-path")
    parser.add_argument("--output")
    return parser.parse_args()


def write_json_atomic(payload, path):
    path.parent.mkdir(parents=True, exist_ok=True)
    tmp_path = path.with_name(f".{path.name}.tmp")
    tmp_path.write_text(json.dumps(payload, indent=2) + "\n")
    tmp_path.replace(path)


def write_csv_atomic(df, path):
    path.parent.mkdir(parents=True, exist_ok=True)
    tmp_path = path.with_name(f".{path.name}.tmp")
    df.to_csv(tmp_path, index=False)
    tmp_path.replace(path)


def write_workbook_atomic(workbook, path):
    path.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False, dir=path.parent) as tmp_file:
        tmp_path = Path(tmp_file.name)
    try:
        workbook.save(tmp_path)
        tmp_path.replace(path)
    finally:
        if tmp_path.exists():
            tmp_path.unlink()


def extract_pdf_layout_text(pdf_bytes):
    require_pdftotext()
    with tempfile.NamedTemporaryFile(suffix=".pdf") as pdf_file:
        pdf_file.write(pdf_bytes)
        pdf_file.flush()
        return subprocess.check_output(["pdftotext", "-layout", pdf_file.name, "-"]).decode(
            "utf-8", errors="ignore"
        )


def canonical_team_slug(raw_name):
    normalized = _normalize_school_name(raw_name)
    slug = _slugify_school(normalized)
    slug = PLAY_IN_NAME_ALIASES.get(slug, slug)
    slug = BRACKET_TEAM_ALIASES.get(slug, slug)
    return normalized, slug


def build_team_lookup(field_df):
    lookup = {}
    for row in field_df.itertuples(index=False):
        keys = {
            ("slug", row.team_slug),
            ("name", _normalize_school_name(row.bracket_team_name)),
            ("name", _normalize_school_name(row.normalized_team_name)),
            ("name", _normalize_school_name(row.school)),
        }
        for source_name in [row.bracket_team_name, row.normalized_team_name, row.school]:
            normalized, slug = canonical_team_slug(source_name)
            keys.add(("name", normalized))
            keys.add(("slug", slug))
        for key in keys:
            lookup[key] = {
                "team_slug": row.team_slug,
                "team_name": row.school,
                "seed": int(row.seed),
                "field_order": int(row.field_order),
            }
    return lookup


def resolve_team(raw_name, expected_seed, team_lookup):
    normalized, slug = canonical_team_slug(raw_name)
    candidate = team_lookup.get(("slug", slug)) or team_lookup.get(("name", normalized))
    if candidate is None:
        raise ValueError(f"Could not map bracket team '{raw_name}'")
    if int(candidate["seed"]) != int(expected_seed):
        raise ValueError(
            f"Bracket team '{raw_name}' mapped to seed {candidate['seed']}, expected seed {expected_seed}"
        )
    return candidate


def parse_seed_slot_line(line):
    line = re.sub(r"^\d{1,2}/\d{1,2}\s+", "", line).strip()
    match = SEED_LINE_RE.match(line)
    if not match:
        fallback_matches = list(TEAM_WITH_RECORD_RE.finditer(line))
        if not fallback_matches:
            return None
        match = fallback_matches[-1]
        seed = int(match.group("seed"))
        body = match.group("body").strip()
    else:
        seed = int(match.group("seed"))
        body = match.group("body").strip()

    record_match = RECORD_IN_LINE_RE.match(body)
    if record_match:
        body = record_match.group("team").strip()
    if not body:
        return None

    play_in_match = PLAY_IN_RE.match(body)
    if play_in_match:
        left = re.sub(rf"^(?:{seed})\s*", "", play_in_match.group("left")).strip()
        right = re.sub(rf"^(?:{seed})\s*", "", play_in_match.group("right")).strip()
        return {
            "kind": "play_in",
            "seed": seed,
            "teams": [left, right],
            "slot_name": f"{seed}-seed play-in",
        }

    return {
        "kind": "team",
        "seed": seed,
        "team_name": body,
        "slot_name": f"{seed} seed",
    }


def parse_region_blocks(pdf_text):
    raw_lines = pdf_text.splitlines()
    midpoint = max(len(line) for line in raw_lines) // 2

    def parse_region_column(column_lines, ordered_regions):
        collected_entries = []
        started = False

        for raw_line in column_lines:
            line = " ".join(raw_line.strip().split())
            if not line:
                continue

            entry_candidates = []
            region_match = re.search(r"\b(EAST|WEST|SOUTH|MIDWEST)\b", line)
            if region_match and region_match.group(1) in ordered_regions:
                before = line[: region_match.start()].strip()
                after = line[region_match.end() :].strip()
                entry_candidates.extend([before, after])
            else:
                entry_candidates.append(line)

            for candidate_line in entry_candidates:
                entry = parse_seed_slot_line(candidate_line)
                if entry is None:
                    continue
                if not started:
                    if entry["kind"] == "team" and entry["seed"] == 1:
                        started = True
                    else:
                        continue

                collected_entries.append(entry)
                if len(collected_entries) == 32:
                    break
            if len(collected_entries) == 32:
                break

        expected_entries = 16 * len(ordered_regions)
        if len(collected_entries) != expected_entries:
            raise ValueError(
                f"Expected {expected_entries} slot entries for regions {ordered_regions}, found {len(collected_entries)}"
            )

        return {
            region: collected_entries[index * 16 : (index + 1) * 16]
            for index, region in enumerate(ordered_regions)
        }

    left_lines = [line[:midpoint] for line in raw_lines]
    right_lines = [line[midpoint:] for line in raw_lines]
    blocks_by_region = {}
    blocks_by_region.update(parse_region_column(left_lines, ["EAST", "SOUTH"]))
    blocks_by_region.update(parse_region_column(right_lines, ["WEST", "MIDWEST"]))

    missing = [region for region in REGION_NAMES if region not in blocks_by_region]
    if missing:
        raise ValueError(f"Missing region blocks: {missing}")
    return blocks_by_region


def build_bracket_games(region_blocks, field_df):
    team_lookup = build_team_lookup(field_df)
    games = []

    for region in REGION_NAMES:
        entries = region_blocks[region]
        seeds = [entry["seed"] for entry in entries]
        if seeds != SEED_SLOT_ORDER:
            raise ValueError(f"{region} seed order mismatch: {seeds}")

        slot_sources = []
        first_four_index = 1
        for entry in entries:
            if entry["kind"] == "team":
                team = resolve_team(entry["team_name"], entry["seed"], team_lookup)
                slot_sources.append(
                    {
                        "type": "team",
                        "id": team["team_slug"],
                    }
                )
                continue

            team_a = resolve_team(entry["teams"][0], entry["seed"], team_lookup)
            team_b = resolve_team(entry["teams"][1], entry["seed"], team_lookup)
            game_id = f"{region.lower()}_ff_{first_four_index}"
            first_four_index += 1
            games.append(
                {
                    "game_id": game_id,
                    "round_order": 0,
                    "round_name": "First Four",
                    "region": region,
                    "slot_label": entry["slot_name"],
                    "source_top_type": "team",
                    "source_top_id": team_a["team_slug"],
                    "source_bottom_type": "team",
                    "source_bottom_id": team_b["team_slug"],
                }
            )
            slot_sources.append({"type": "game", "id": game_id})

        round_sources = slot_sources
        round_specs = [
            ("Round of 64", "r64", 8),
            ("Round of 32", "r32", 4),
            ("Sweet 16", "s16", 2),
            ("Elite 8", "e8", 1),
        ]
        for round_order, (round_name, prefix, game_count) in enumerate(round_specs, start=1):
            next_sources = []
            for game_index in range(game_count):
                source_top = round_sources[game_index * 2]
                source_bottom = round_sources[game_index * 2 + 1]
                game_id = f"{region.lower()}_{prefix}_{game_index + 1}"
                games.append(
                    {
                        "game_id": game_id,
                        "round_order": round_order,
                        "round_name": round_name,
                        "region": region,
                        "slot_label": f"{round_name} {game_index + 1}",
                        "source_top_type": source_top["type"],
                        "source_top_id": source_top["id"],
                        "source_bottom_type": source_bottom["type"],
                        "source_bottom_id": source_bottom["id"],
                    }
                )
                next_sources.append({"type": "game", "id": game_id})
            round_sources = next_sources

    for semifinal_index, (left_region, right_region) in enumerate(SEMIFINAL_PAIRINGS, start=1):
        games.append(
            {
                "game_id": f"final_four_{semifinal_index}",
                "round_order": 5,
                "round_name": "Final Four",
                "region": f"{left_region}/{right_region}",
                "slot_label": f"{left_region} vs {right_region}",
                "source_top_type": "game",
                "source_top_id": f"{left_region.lower()}_e8_1",
                "source_bottom_type": "game",
                "source_bottom_id": f"{right_region.lower()}_e8_1",
            }
        )

    games.append(
        {
            "game_id": "championship",
            "round_order": 6,
            "round_name": "Championship",
            "region": "National",
            "slot_label": "National Championship",
            "source_top_type": "game",
            "source_top_id": "final_four_1",
            "source_bottom_type": "game",
            "source_bottom_id": "final_four_2",
        }
    )

    return games


def load_probability_lookup(matrix_path):
    matrix_df = pd.read_csv(matrix_path)
    matrix_df = matrix_df.set_index("team_slug")
    return matrix_df


def choose_winner(team_a, team_b, prob_a):
    if prob_a > 0.5:
        return team_a["team_slug"], prob_a
    if prob_a < 0.5:
        return team_b["team_slug"], 1.0 - prob_a
    if int(team_a["seed"]) != int(team_b["seed"]):
        if int(team_a["seed"]) < int(team_b["seed"]):
            return team_a["team_slug"], prob_a
        return team_b["team_slug"], 1.0 - prob_a
    if int(team_a["field_order"]) <= int(team_b["field_order"]):
        return team_a["team_slug"], prob_a
    return team_b["team_slug"], 1.0 - prob_a


def simulate_bracket(games, field_df, probability_lookup):
    team_meta = {
        row.team_slug: {
            "team_slug": row.team_slug,
            "team_name": row.school,
            "seed": int(row.seed),
            "field_order": int(row.field_order),
        }
        for row in field_df.itertuples(index=False)
    }
    winners = {}
    game_rows = []
    game_results = {}

    def resolve_source(source_type, source_id):
        if source_type == "team":
            return team_meta[source_id]
        winner_slug = winners.get(source_id)
        if winner_slug is None:
            raise ValueError(f"Missing winner for upstream game {source_id}")
        return team_meta[winner_slug]

    for game in sorted(games, key=lambda item: (item["round_order"], item["game_id"])):
        team_a = resolve_source(game["source_top_type"], game["source_top_id"])
        team_b = resolve_source(game["source_bottom_type"], game["source_bottom_id"])
        prob_a = float(probability_lookup.loc[team_a["team_slug"], team_b["team_slug"]])
        prob_b = 1.0 - prob_a
        winner_slug, winner_prob = choose_winner(team_a, team_b, prob_a)
        winners[game["game_id"]] = winner_slug
        winner = team_meta[winner_slug]
        loser = team_b if winner_slug == team_a["team_slug"] else team_a
        game_result = {
            **game,
            "team_a_slug": team_a["team_slug"],
            "team_a_name": team_a["team_name"],
            "seed_a": team_a["seed"],
            "team_b_slug": team_b["team_slug"],
            "team_b_name": team_b["team_name"],
            "seed_b": team_b["seed"],
            "win_prob_a": prob_a,
            "win_prob_b": prob_b,
            "picked_winner_slug": winner["team_slug"],
            "picked_winner_name": winner["team_name"],
            "picked_seed": winner["seed"],
            "picked_winner_prob": winner_prob,
            "picked_loser_slug": loser["team_slug"],
            "picked_loser_name": loser["team_name"],
            "picked_loser_seed": loser["seed"],
        }
        game_rows.append(game_result)
        game_results[game["game_id"]] = game_result

    expected_games = 67
    if len(game_rows) != expected_games:
        raise ValueError(f"Expected {expected_games} simulated games, found {len(game_rows)}")
    return pd.DataFrame(game_rows), game_results


def set_region_columns(ws, left_side):
    if left_side:
        return {
            "r64_team": 1,
            "r64_prob": 2,
            "r32_team": 4,
            "r32_prob": 5,
            "s16_team": 7,
            "s16_prob": 8,
            "e8_team": 10,
            "e8_prob": 11,
            "semi_team": 13,
            "semi_prob": 14,
        }
    return {
        "semi_team": 18,
        "semi_prob": 19,
        "e8_team": 21,
        "e8_prob": 22,
        "s16_team": 24,
        "s16_prob": 25,
        "r32_team": 27,
        "r32_prob": 28,
        "r64_team": 30,
        "r64_prob": 31,
    }


def round_rows(start_row):
    return {
        "Round of 64": [start_row + offset for offset in [0, 4, 8, 12, 16, 20, 24, 28]],
        "Round of 32": [start_row + offset for offset in [2, 10, 18, 26]],
        "Sweet 16": [start_row + offset for offset in [6, 22]],
        "Elite 8": [start_row + 14],
    }


def style_team_cell(cell, is_winner):
    cell.border = CELL_BORDER
    cell.alignment = Alignment(horizontal="left", vertical="center")
    if is_winner:
        cell.fill = WINNER_FILL
        cell.font = Font(bold=True)
    else:
        cell.fill = LOSER_FILL
        cell.font = Font(bold=False)


def style_prob_cell(cell, is_winner):
    cell.border = CELL_BORDER
    cell.alignment = Alignment(horizontal="center", vertical="center")
    cell.number_format = "0.0%"
    if is_winner:
        cell.fill = WINNER_FILL
        cell.font = Font(bold=True)
    else:
        cell.fill = LOSER_FILL


def write_matchup(ws, row, team_col, prob_col, game_result):
    top_is_winner = game_result["picked_winner_slug"] == game_result["team_a_slug"]

    ws.cell(row=row, column=team_col, value=f"({game_result['seed_a']}) {game_result['team_a_name']}")
    style_team_cell(ws.cell(row=row, column=team_col), top_is_winner)
    ws.cell(row=row, column=prob_col, value=game_result["win_prob_a"])
    style_prob_cell(ws.cell(row=row, column=prob_col), top_is_winner)

    ws.cell(row=row + 1, column=team_col, value=f"({game_result['seed_b']}) {game_result['team_b_name']}")
    style_team_cell(ws.cell(row=row + 1, column=team_col), not top_is_winner)
    ws.cell(row=row + 1, column=prob_col, value=game_result["win_prob_b"])
    style_prob_cell(ws.cell(row=row + 1, column=prob_col), not top_is_winner)


def build_bracket_sheet(workbook, game_results, champion_name):
    ws = workbook.active
    ws.title = "Bracket"
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "A4"
    ws["A1"] = "2026 NCAA Tournament Bracket Picks"
    ws["A1"].font = Font(bold=True, color="FFFFFF", size=16)
    ws["A1"].fill = TITLE_FILL
    ws["A1"].alignment = Alignment(horizontal="center")
    ws.merge_cells("A1:AE1")

    ws["A2"] = "Picked strictly from the model probability matrix"
    ws["A2"].font = Font(italic=True)
    ws.merge_cells("A2:AE2")
    ws["O3"] = "Champion"
    ws["O3"].font = Font(bold=True)
    ws["O4"] = champion_name
    ws["O4"].font = Font(bold=True, size=14)
    ws["O4"].fill = FINAL_FILL
    ws["O4"].border = CELL_BORDER
    ws["O4"].alignment = Alignment(horizontal="center")
    ws.merge_cells("O4:Q4")

    for col in range(1, 32):
        width = 5 if col in {2, 5, 8, 11, 14, 19, 22, 25, 28, 31} else 22
        ws.column_dimensions[get_column_letter(col)].width = width

    region_layouts = [
        ("EAST", 4, True),
        ("WEST", 42, True),
        ("SOUTH", 4, False),
        ("MIDWEST", 42, False),
    ]

    game_positions = {}
    for region, start_row, left_side in region_layouts:
        cols = set_region_columns(ws, left_side)
        rows = round_rows(start_row)

        region_header_col_start = cols["r64_team"] if left_side else cols["semi_team"]
        region_header_col_end = cols["semi_prob"] if left_side else cols["r64_prob"]
        ws.cell(row=start_row - 1, column=region_header_col_start, value=region)
        header_cell = ws.cell(row=start_row - 1, column=region_header_col_start)
        header_cell.font = Font(bold=True, color="000000")
        header_cell.fill = PatternFill("solid", fgColor=REGION_COLORS[region])
        header_cell.alignment = Alignment(horizontal="center")
        header_cell.border = CELL_BORDER
        ws.merge_cells(
            start_row=start_row - 1,
            start_column=region_header_col_start,
            end_row=start_row - 1,
            end_column=region_header_col_end,
        )

        for round_name, prefix in [
            ("Round of 64", "r64"),
            ("Round of 32", "r32"),
            ("Sweet 16", "s16"),
            ("Elite 8", "e8"),
        ]:
            team_col = cols[f"{prefix}_team"]
            prob_col = cols[f"{prefix}_prob"]
            for index, row in enumerate(rows[round_name], start=1):
                game_id = f"{region.lower()}_{prefix}_{index}"
                write_matchup(ws, row, team_col, prob_col, game_results[game_id])
                game_positions[game_id] = (row, team_col, prob_col)

    semifinal_rows = [35, 35]
    semifinal_cols = [(13, 14), (18, 19)]
    for index, pair in enumerate(SEMIFINAL_PAIRINGS, start=1):
        row = semifinal_rows[index - 1]
        team_col, prob_col = semifinal_cols[index - 1]
        game_id = f"final_four_{index}"
        ws.cell(row=row - 1, column=team_col, value=f"Final Four: {pair[0]} vs {pair[1]}")
        ws.cell(row=row - 1, column=team_col).font = Font(bold=True)
        ws.merge_cells(start_row=row - 1, start_column=team_col, end_row=row - 1, end_column=prob_col)
        write_matchup(ws, row, team_col, prob_col, game_results[game_id])

    ws["O36"] = "Championship"
    ws["O36"].font = Font(bold=True)
    ws.merge_cells("O36:Q36")
    championship_result = game_results["championship"]
    write_matchup(ws, 37, 15, 16, championship_result)
    ws.cell(row=37, column=17, value="Winner")
    ws.cell(row=37, column=17).font = Font(bold=True)
    ws.cell(row=38, column=17, value=championship_result["picked_winner_name"])
    ws.cell(row=38, column=17).font = Font(bold=True)
    ws.cell(row=38, column=17).fill = FINAL_FILL
    ws.cell(row=38, column=17).border = CELL_BORDER

    ws["M42"] = "First Four"
    ws["M42"].font = Font(bold=True)
    ws["M42"].alignment = Alignment(horizontal="center")
    ws.merge_cells("M42:S42")
    first_four_ids = sorted(game_id for game_id in game_results if "_ff_" in game_id)
    first_four_rows = [44, 48, 52, 56]
    for row, game_id in zip(first_four_rows, first_four_ids):
        result = game_results[game_id]
        ws.cell(row=row - 1, column=13, value=f"{result['region']} {result['slot_label']}")
        ws.cell(row=row - 1, column=13).font = Font(bold=True)
        ws.merge_cells(start_row=row - 1, start_column=13, end_row=row - 1, end_column=16)
        write_matchup(ws, row, 13, 14, result)


def autosize_data_sheet(ws):
    for column_cells in ws.columns:
        values = ["" if cell.value is None else str(cell.value) for cell in column_cells]
        max_length = max(len(value) for value in values) if values else 0
        ws.column_dimensions[column_cells[0].column_letter].width = min(max(max_length + 2, 12), 28)


def build_game_data_sheet(workbook, game_rows):
    ws = workbook.create_sheet("Game Data")
    ws.freeze_panes = "A2"
    ordered_cols = [
        "game_id",
        "round_name",
        "region",
        "slot_label",
        "team_a_name",
        "seed_a",
        "team_b_name",
        "seed_b",
        "win_prob_a",
        "win_prob_b",
        "picked_winner_name",
        "picked_seed",
        "picked_winner_prob",
    ]
    ws.append(ordered_cols)
    for row in game_rows[ordered_cols].itertuples(index=False):
        ws.append(list(row))

    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = TITLE_FILL
        cell.border = CELL_BORDER
        cell.alignment = Alignment(horizontal="center")

    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.border = CELL_BORDER
            if cell.column in {9, 10, 13}:
                cell.number_format = "0.0%"
    autosize_data_sheet(ws)


def verify_workbook(path):
    workbook = load_workbook(path, data_only=True)
    if "Bracket" not in workbook.sheetnames or "Game Data" not in workbook.sheetnames:
        raise ValueError("Workbook verification failed: missing required sheets")
    if workbook["Bracket"]["O4"].value is None:
        raise ValueError("Workbook verification failed: champion cell is blank")


def main():
    args = parse_args()
    tournament_dir = Path(args.tournament_dir) if args.tournament_dir else DEFAULT_TOURNAMENT_DIR / str(args.season_year)
    probability_matrix_path = (
        Path(args.probability_matrix_path) if args.probability_matrix_path else tournament_dir / "probability_matrix.csv"
    )
    team_field_path = Path(args.team_field_path) if args.team_field_path else tournament_dir / "team_field.csv"
    output_path = Path(args.output) if args.output else DEFAULT_OUTPUT_DIR / str(args.season_year) / "bracket.xlsx"
    bracket_games_path = output_path.parent / "bracket_games.csv"
    run_meta_path = output_path.parent / "run_meta.json"

    field_source_url = args.field_source_url or default_field_source_url(args.season_year)
    field_df = pd.read_csv(team_field_path)
    probability_lookup = load_probability_lookup(probability_matrix_path)

    pdf_bytes, resolved_url = fetch_bracket_pdf(field_source_url)
    pdf_text = extract_pdf_layout_text(pdf_bytes)
    region_blocks = parse_region_blocks(pdf_text)
    games = build_bracket_games(region_blocks, field_df)
    game_rows, game_results = simulate_bracket(games, field_df, probability_lookup)

    champion_name = game_results["championship"]["picked_winner_name"]
    workbook = Workbook()
    build_bracket_sheet(workbook, game_results, champion_name)
    build_game_data_sheet(workbook, game_rows)
    write_workbook_atomic(workbook, output_path)
    verify_workbook(output_path)

    write_csv_atomic(game_rows, bracket_games_path)
    write_json_atomic(
        {
            "season_year": args.season_year,
            "cutoff_date": args.cutoff_date,
            "field_source_url": field_source_url,
            "resolved_field_source_url": resolved_url,
            "team_field_path": str(team_field_path),
            "probability_matrix_path": str(probability_matrix_path),
            "output_path": str(output_path),
            "game_rows": int(len(game_rows)),
            "champion": champion_name,
        },
        run_meta_path,
    )

    print(f"Saved workbook to {output_path}")
    print(f"Saved game data to {bracket_games_path}")
    print(f"Champion: {champion_name}")


if __name__ == "__main__":
    main()
